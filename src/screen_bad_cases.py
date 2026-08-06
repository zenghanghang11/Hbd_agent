# -*- coding: utf-8 -*-
"""地址 bad case 筛查脚本.

用法:
    python3 src/screen_bad_cases.py \
        --input data/addresses.txt \
        --outdir output

产出:
    output/bad_case_screening.csv   逐条明细(UTF-8-BOM, Excel 可直接打开)
    output/bad_case_screening.xlsx  三个工作表: 明细 / 规则统计 / 数据集级发现
    output/bad_case_report.md       人读版筛查报告

规则分三档:
    严重(P0) 无法可靠定位到唯一楼宇, 必须人工修正
    一般(P1) 影响去重与一致性, 建议修正
    轻微(P2) 纯格式问题, 脚本可自动修复
"""

import argparse
import csv
import json
import os
import re
import sys
from collections import Counter, defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import hk_address as hk

# --------------------------------------------------------------------------
# 规则定义: code -> (类别, 严重度, 说明, 建议处理)
# --------------------------------------------------------------------------

SEVERITY_ORDER = {"严重": 3, "一般": 2, "轻微": 1, "": 0}

RULES = {
    # --- 格式类 (轻微, 脚本可自动修复) ---
    "F001": ("格式", "轻微", "使用全角标点分隔", "统一替换为半角逗号"),
    "F002": ("格式", "轻微", "首尾存在多余空白", "trim 后入库"),
    "F003": ("格式", "轻微", "座号英文字母小写", "统一大写, 如 e座 -> E座"),
    "F004": ("格式", "轻微", "含全角字母/数字", "NFKC 归一化为半角"),
    "F005": ("格式", "轻微", "存在空的地址字段段", "去除空段"),
    "F006": ("格式", "轻微", "无分隔符连写, 依赖模糊切分", "按 座,屋苑,街道,地區,大區 补分隔符"),
    "F007": ("格式", "轻微", "期数使用罗马数字", "转写为阿拉伯数字, 如 XIIB期 -> 12B期"),

    # --- 重复类 (一般) ---
    "D001": ("重复", "一般", "与其他记录原文完全重复", "去重, 保留一条"),
    "D002": ("重复", "一般", "归一化后与其他记录重复(仅格式差异)", "归一化后去重"),
    "D003": ("重复", "一般", "同一楼宇存在详略两种写法", "统一采用信息最全的写法"),

    # --- 完整性 (严重/一般) ---
    "C001": ("完整性", "严重", "缺少大區(九龍/新界/港島)", "补全大區; 地區已知时可反查补全"),
    "C002": ("完整性", "严重", "缺少地區", "补全地區"),
    "C003": ("完整性", "严重", "只有座/閣名, 缺少屋苑, 无法唯一定位", "补全屋苑名称"),
    "C004": ("完整性", "一般", "缺少街道", "补全街道, 便于地图匹配"),
    "C005": ("完整性", "一般", "缺少座/期号", "确认是整个屋苑还是漏填座号"),

    # --- 内容异常 (严重/一般) ---
    "A001": ("异常", "严重", "尾部粘连了表单字段名等噪声内容", "剔除噪声 token 后重新入库"),
    "A002": ("异常", "严重", "疑似形近错别字", "核实用字后修正"),
    "A003": ("异常", "一般", "座号数值偏大, 疑似超出该屋苑实际座数", "对照屋苑座数表核实"),
    "A004": ("异常", "一般", "存在无法归类的地址片段", "人工确认该片段含义"),
    "A005": ("异常", "一般", "结构解析置信度低", "补分隔符后重跑解析"),
    "A006": ("异常", "一般", "可依据同批数据中的同族记录补全", "按同族记录补全屋苑/街道/地區/大區"),
}

# 座号数值超过该阈值即提示核实
HIGH_BLOCK_THRESHOLD = 20

CN_DIGIT = {"零": 0, "一": 1, "二": 2, "兩": 2, "两": 2, "三": 3, "四": 4,
            "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}

ROMAN = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}


def cn_to_int(s):
    """中文数字转阿拉伯数字, 支持 十/二十/三十/十七 这类写法."""
    s = re.sub(r"[^" + "".join(CN_DIGIT) + "十百]", "", s)
    if not s:
        return None
    if "百" in s or "十" in s:
        total, unit_seen = 0, False
        num = 0
        for ch in s:
            if ch == "十":
                total += (num or 1) * 10
                num, unit_seen = 0, True
            elif ch == "百":
                total += (num or 1) * 100
                num, unit_seen = 0, True
            else:
                num = CN_DIGIT.get(ch, 0)
        return total + num if unit_seen else num
    return CN_DIGIT.get(s[0]) if len(s) == 1 else int("".join(str(CN_DIGIT.get(c, 0)) for c in s))


def roman_to_int(s):
    total, prev = 0, 0
    for ch in reversed(s.upper()):
        v = ROMAN.get(ch, 0)
        total = total - v if v < prev else total + v
        prev = max(prev, v)
    return total or None


def block_number(block):
    """从座号中抽出数值, 抽不出返回 None."""
    if not block:
        return None
    m = re.search(r"\d+", block)
    if m:
        return int(m.group(0))
    m = re.search(r"第([" + "".join(CN_DIGIT) + "十百]+)座", block)
    if m:
        return cn_to_int(m.group(1))
    return None


# --------------------------------------------------------------------------
# 记录
# --------------------------------------------------------------------------

class Record:
    def __init__(self, idx, raw):
        self.idx = idx
        self.raw = raw.rstrip("\n")
        self.norm, fmt_issues = hk.normalize(self.raw)
        self.parsed = hk.parse(self.norm)
        self.canonical = hk.canonical(self.parsed)
        self.key = hk.dedup_key(self.parsed, self.norm)
        self.hits = []          # [(code, 细节)]
        self.dup_group = ""

        code_of = {"FULLWIDTH_PUNCT": "F001", "TRAILING_WHITESPACE": "F002",
                   "LOWERCASE_BLOCK_LETTER": "F003", "FULLWIDTH_ALNUM": "F004",
                   "EMPTY_SEGMENT": "F005"}
        for issue in fmt_issues:
            if issue in code_of:
                self.add(code_of[issue])

    def add(self, code, detail="", severity=None):
        """severity 传入时覆盖规则默认严重度(同一规则在不同上下文危害不同)."""
        if code not in self.codes:
            self.hits.append((code, detail, severity or RULES[code][1]))

    @property
    def codes(self):
        return [h[0] for h in self.hits]

    @property
    def max_severity(self):
        if not self.hits:
            return ""
        return max((h[2] for h in self.hits), key=lambda s: SEVERITY_ORDER[s])

    @property
    def is_bad_case(self):
        return SEVERITY_ORDER[self.max_severity] >= SEVERITY_ORDER["一般"]


# --------------------------------------------------------------------------
# 逐条规则
# --------------------------------------------------------------------------

def check_row_rules(rec):
    p = rec.parsed

    # 无分隔符连写
    if "," not in rec.norm and len(rec.norm) > 6:
        rec.add("F006")
    # 罗马数字期数
    if re.search(r"[IVXLCDM]{2,}[A-Za-z]?期", rec.norm):
        m = re.search(r"([IVXLCDM]{2,})([A-Za-z]?)期", rec.norm)
        n = roman_to_int(m.group(1))
        rec.add("F007", "%s%s期 -> %s%s期" % (m.group(1), m.group(2), n, m.group(2)))

    # 粘连的噪声 token(表单字段名等)
    if p["noise"]:
        rec.add("A001", "「%s%s」中的「%s」为无效内容" % (p["region"], p["noise"], p["noise"]))
    else:
        for tok in hk.NOISE_TOKENS:
            if rec.norm.endswith(tok) and rec.norm != tok:
                rec.add("A001", "尾部噪声: 「%s」" % tok)
                break

    # 形近错字
    for bad, good in hk.SUSPECT_CHAR_MAP.items():
        if re.search(r"[一-鿿]{1,3}%s(,|$)" % bad, rec.norm):
            m = re.search(r"([一-鿿]{1,3})%s" % bad, rec.norm)
            rec.add("A002", "「%s%s」疑为「%s%s」" % (m.group(1), bad, m.group(1), good))
            break

    # 完整性
    if not p["region"] or not hk.is_region(p["region"]):
        if p["district"]:
            rec.add("C001", "地區「%s」已知, 大區可反查补全" % p["district"], severity="一般")
        else:
            rec.add("C001")
    if not p["district"]:
        rec.add("C002")
    if not p["estate"]:
        rec.add("C003")
    if not p["street"]:
        rec.add("C004")
    if not p["block"] and not p["phase"]:
        rec.add("C005")

    # 座号数值偏大
    n = block_number(p["block"])
    if n and n >= HIGH_BLOCK_THRESHOLD:
        rec.add("A003", "座号=%d" % n)

    # 未识别片段 / 低置信度
    if p["leftover"]:
        rec.add("A004", "未归类: %s" % "、".join(p["leftover"]))
    if p["confidence"] == "low":
        rec.add("A005")


# --------------------------------------------------------------------------
# 跨条规则(重复 / 同族补全)
# --------------------------------------------------------------------------

def check_cross_rules(records):
    # 1. 原文完全重复
    raw_groups = defaultdict(list)
    for r in records:
        raw_groups[r.raw].append(r)
    for raw, group in raw_groups.items():
        if len(group) > 1:
            gid = "R%02d" % group[0].idx
            for r in group:
                r.add("D001", "与第 %s 条重复" % ",".join(str(x.idx) for x in group if x is not r))
                r.dup_group = gid

    # 2. 归一化后重复
    norm_groups = defaultdict(list)
    for r in records:
        norm_groups[r.norm].append(r)
    for norm, group in norm_groups.items():
        if len(group) > 1 and len({r.raw for r in group}) > 1:
            gid = "N%02d" % group[0].idx
            for r in group:
                r.add("D002", "与第 %s 条归一化后一致" % ",".join(str(x.idx) for x in group if x is not r))
                r.dup_group = r.dup_group or gid

    # 3. 同一楼宇的详略两种写法(按楼宇名分组, 屋苑不冲突才算同一栋)
    key_groups = defaultdict(list)
    for r in records:
        b = hk.building_name(r.parsed)
        key_groups[b if b else r.key].append(r)
    for key, group in key_groups.items():
        estates = {r.parsed["estate"] for r in group if r.parsed["estate"]}
        estates -= {key}
        if len(estates) > 1:
            continue  # 同名楼宇分属不同屋苑, 不能当作重复
        if len(group) > 1 and len({r.norm for r in group}) > 1:
            fullest = max(group, key=lambda r: len(r.norm))
            gid = "S%02d" % group[0].idx
            for r in group:
                if r.norm != fullest.norm:
                    r.add("D003", "与第 %d 条同楼宇, 建议统一为「%s」" % (fullest.idx, fullest.norm))
                r.dup_group = r.dup_group or gid

    # 4. 同族补全: 同一"字头 + 閣/樓"家族里, 有完整地址的记录可用于补全残缺记录
    family = defaultdict(list)
    for r in records:
        blk = r.parsed["block"] or r.parsed["estate"]
        m = re.fullmatch(r"([一-鿿])([一-鿿]{1,2})([閣樓阁楼])", blk or "")
        if m:
            family[(m.group(1), m.group(3))].append(r)
    for (prefix, tail), group in family.items():
        donors = [r for r in group if r.parsed["estate"] and r.parsed["region"]
                  and r.parsed["block"]]
        needy = [r for r in group if not r.parsed["estate"] or not r.parsed["region"]]
        if not donors or not needy:
            continue
        donor = max(donors, key=lambda r: len(r.norm))
        for r in needy:
            suggestion = ",".join(x for x in [r.parsed["block"] or r.parsed["estate"],
                                              donor.parsed["estate"],
                                              donor.parsed["street"],
                                              donor.parsed["district"],
                                              donor.parsed["region"]] if x)
            r.add("A006", "同族参考第 %d 条, 建议补全为「%s」" % (donor.idx, suggestion))


def dataset_findings(records):
    """数据集级观察(不判定单条 bad case, 供人工复核)."""
    out = []

    total = len(records)
    bad = sum(1 for r in records if r.is_bad_case)
    uniq = len({r.norm for r in records})
    out.append(("记录规模", "共 %d 条, 归一化后唯一 %d 条, 冗余 %d 条(%.0f%%)"
                % (total, uniq, total - uniq, (total - uniq) * 100.0 / total)))
    out.append(("bad case 占比", "%d / %d = %.0f%%" % (bad, total, bad * 100.0 / total)))

    # 同屋苑座号连续性
    by_estate = defaultdict(set)
    for r in records:
        est = (r.parsed["estate"] or "") + (r.parsed["phase"] or "")
        n = block_number(r.parsed["block"])
        if est and n:
            by_estate[est].add(n)
    for est, nums in sorted(by_estate.items()):
        if len(nums) >= 3:
            missing = sorted(set(range(min(nums), max(nums) + 1)) - nums)
            if missing:
                out.append(("座号不连续",
                            "%s 出现 %s 座, 缺 %s 座(香港常见跳过 4/13, 需核实是漏采还是本就没有)"
                            % (est, "/".join(map(str, sorted(nums))), "/".join(map(str, missing)))))

    # 座号写法混用
    styles = Counter()
    for r in records:
        b = r.parsed["block"]
        if not b:
            continue
        if re.search(r"第[一二三四五六七八九十]", b):
            styles["中文数字(第七座)"] += 1
        elif re.search(r"\d", b):
            styles["阿拉伯数字(第1座)"] += 1
        elif re.search(r"[A-Z]", b):
            styles["英文字母(B座)"] += 1
        else:
            styles["楼宇名(愉澤閣)"] += 1
    if len(styles) > 1:
        out.append(("座号写法混用",
                    "同一批数据混用 %d 种写法: %s" %
                    (len(styles), ", ".join("%s×%d" % kv for kv in styles.most_common()))))

    # 地址粒度混用
    granularity = Counter("整苑/无座号" if not (r.parsed["block"] or r.parsed["phase"])
                          else "到座" for r in records)
    if len(granularity) > 1:
        out.append(("地址粒度混用",
                    ", ".join("%s×%d" % kv for kv in granularity.most_common())))
    return out


# --------------------------------------------------------------------------
# 输出
# --------------------------------------------------------------------------

HEADERS = ["序号", "原始地址", "归一化地址", "规范地址", "大區", "地區", "地區(上级)",
           "街道", "屋苑", "期", "座", "未识别片段", "解析置信度", "是否bad case",
           "最高严重度", "命中规则数", "命中规则", "问题描述", "建议处理", "重复组"]


def row_of(rec):
    p = rec.parsed
    descs, sugs = [], []
    for code, detail, sev in sorted(rec.hits, key=lambda h: -SEVERITY_ORDER[h[2]]):
        cat, _default_sev, desc, sug = RULES[code]
        descs.append("[%s·%s]%s%s" % (code, sev, desc, ("(%s)" % detail) if detail else ""))
        sugs.append(sug)
    return [
        rec.idx, rec.raw, rec.norm, rec.canonical,
        p["region"], p["district"], p["district_parent"],
        p["street"], p["estate"], p["phase"], p["block"],
        "、".join(p["leftover"]), {"high": "高", "low": "低"}[p["confidence"]],
        "是" if rec.is_bad_case else "否",
        rec.max_severity or "-",
        len(rec.hits),
        ";".join(rec.codes) or "-",
        " ".join(descs) or "无问题",
        "; ".join(dict.fromkeys(sugs)) or "-",
        rec.dup_group or "-",
    ]


def write_csv(records, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        for r in records:
            w.writerow(row_of(r))


def write_rule_stats_rows(records):
    counter = Counter()
    for r in records:
        counter.update(r.codes)
    rows = []
    for code in sorted(RULES, key=lambda c: (-SEVERITY_ORDER[RULES[c][1]], c)):
        cat, sev, desc, sug = RULES[code]
        n = counter.get(code, 0)
        actual = {h[2] for r in records for h in r.hits if h[0] == code}
        sev = "/".join(sorted(actual, key=lambda x: -SEVERITY_ORDER[x])) or sev
        rows.append([code, cat, sev, desc, n,
                     ",".join(str(r.idx) for r in records if code in r.codes) or "-",
                     sug])
    return rows


def write_xlsx(records, findings, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[warn] 未安装 openpyxl, 跳过 xlsx 输出 (pip install openpyxl)")
        return False

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    sev_fill = {"严重": PatternFill("solid", fgColor="FFC7CE"),
                "一般": PatternFill("solid", fgColor="FFEB9C"),
                "轻微": PatternFill("solid", fgColor="E2EFDA")}

    ws = wb.active
    ws.title = "明细"
    ws.append(HEADERS)
    for r in records:
        ws.append(row_of(r))
    for i, cell in enumerate(ws[1], 1):
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sev_col = HEADERS.index("最高严重度") + 1
    for row in ws.iter_rows(min_row=2):
        fill = sev_fill.get(row[sev_col - 1].value)
        if fill:
            row[sev_col - 1].fill = fill
    widths = {"原始地址": 34, "归一化地址": 34, "规范地址": 36, "问题描述": 60,
              "建议处理": 40, "屋苑": 14, "街道": 12, "地區": 10, "地區(上级)": 11,
              "大區": 10}
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 11)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("规则统计")
    ws2.append(["规则代码", "类别", "严重度", "规则说明", "命中条数", "命中序号", "建议处理"])
    for row in write_rule_stats_rows(records):
        ws2.append(row)
    for cell in ws2[1]:
        cell.font, cell.fill = head_font, head_fill
    for i, w in enumerate([12, 10, 10, 40, 10, 26, 40], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("数据集级发现")
    ws3.append(["观察项", "结论"])
    for k, v in findings:
        ws3.append([k, v])
    for cell in ws3[1]:
        cell.font, cell.fill = head_font, head_fill
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 90

    wb.save(path)
    return True


def write_report(records, findings, path):
    total = len(records)
    bad = [r for r in records if r.is_bad_case]
    by_sev = Counter(r.max_severity or "无问题" for r in records)
    counter = Counter()
    for r in records:
        counter.update(r.codes)

    L = []
    L.append("# 地址 bad case 筛查报告\n")
    L.append("- 输入记录: **%d** 条" % total)
    L.append("- 判定为 bad case: **%d** 条 (%.0f%%)" % (len(bad), len(bad) * 100.0 / total))
    L.append("- 严重度分布: " + ", ".join("%s %d 条" % (k, v) for k, v in by_sev.most_common()))
    L.append("\n> 判定口径: 命中任一「严重」或「一般」规则即计为 bad case;"
             "仅命中「轻微」规则的记录可由脚本自动修复, 不计入.\n")

    L.append("## 一、规则命中统计\n")
    L.append("| 规则 | 类别 | 严重度 | 说明 | 命中 | 命中序号 |")
    L.append("|---|---|---|---|---|---|")
    for code, cat, sev, desc, n, idxs, sug in write_rule_stats_rows(records):
        if n:
            L.append("| %s | %s | %s | %s | %d | %s |" % (code, cat, sev, desc, n, idxs))

    L.append("\n## 二、严重(P0)问题清单\n")
    p0 = [r for r in records if r.max_severity == "严重"]
    if p0:
        L.append("| 序号 | 原始地址 | 问题 | 建议 |")
        L.append("|---|---|---|---|")
        for r in p0:
            row = row_of(r)
            L.append("| %d | `%s` | %s | %s |" % (
                r.idx, r.raw, row[HEADERS.index("问题描述")], row[HEADERS.index("建议处理")]))
    else:
        L.append("无\n")

    L.append("\n## 三、一般(P1)问题清单\n")
    p1 = [r for r in records if r.max_severity == "一般"]
    if p1:
        L.append("| 序号 | 原始地址 | 问题 | 建议 |")
        L.append("|---|---|---|---|")
        for r in p1:
            row = row_of(r)
            L.append("| %d | `%s` | %s | %s |" % (
                r.idx, r.raw, row[HEADERS.index("问题描述")], row[HEADERS.index("建议处理")]))
    else:
        L.append("无\n")

    L.append("\n## 四、数据集级发现\n")
    for k, v in findings:
        L.append("- **%s**: %s" % (k, v))

    L.append("\n## 五、建议的处理优先级\n")
    L.append("1. **P0 先修**: 缺屋苑/缺地區大區、噪声粘连、疑似错别字 —— 这些无法唯一定位到楼宇, "
             "直接影响下游匹配成功率.")
    L.append("2. **P1 再修**: 重复记录与详略不一 —— 先归一化再去重, 可直接压缩数据量.")
    L.append("3. **P2 自动化**: 全角标点、大小写、空白、罗马数字期数 —— 入库前统一走 "
             "`hk_address.normalize()`, 无需人工.")
    L.append("4. **入库校验**: 建议把本脚本的 P0 规则做成写入前的拦截校验, 从源头阻断.\n")

    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


# --------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="地址 bad case 筛查")
    ap.add_argument("--input", default="data/addresses.txt", help="地址列表, 每行一条")
    ap.add_argument("--outdir", default="output", help="输出目录")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8") as f:
        lines = [ln.rstrip("\n") for ln in f if ln.strip()]

    records = [Record(i, ln) for i, ln in enumerate(lines, 1)]
    for r in records:
        check_row_rules(r)
    check_cross_rules(records)
    findings = dataset_findings(records)

    os.makedirs(args.outdir, exist_ok=True)
    csv_path = os.path.join(args.outdir, "bad_case_screening.csv")
    xlsx_path = os.path.join(args.outdir, "bad_case_screening.xlsx")
    md_path = os.path.join(args.outdir, "bad_case_report.md")
    json_path = os.path.join(args.outdir, "bad_case_screening.json")

    write_csv(records, csv_path)
    ok_xlsx = write_xlsx(records, findings, xlsx_path)
    write_report(records, findings, md_path)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump([dict(zip(HEADERS, row_of(r))) for r in records],
                  f, ensure_ascii=False, indent=2)

    bad = sum(1 for r in records if r.is_bad_case)
    print("共 %d 条, bad case %d 条 (%.0f%%)" % (len(records), bad, bad * 100.0 / len(records)))
    print("已输出:")
    print("  " + csv_path)
    if ok_xlsx:
        print("  " + xlsx_path)
    print("  " + md_path)
    print("  " + json_path)


if __name__ == "__main__":
    main()
