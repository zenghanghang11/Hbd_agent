# -*- coding: utf-8 -*-
"""对地址列表批量调 CMHK 接口, 提取 recogFlag=Y 的地址及其房间号与楼层.

三种运行模式:
    --mode live   真机调用接口(需要 curl 模板, 见下). 产出真实结果.
    --mode mock   用 tests/fixtures 下的构造样例跑通全流程, 验证提取逻辑, 不发网络请求.
    --mode plan   只产出查询计划(输入地址 -> 查询关键词 + 离线筛查结论), 不发网络请求.

live 模式准备:
    1. 浏览器打开 https://www.hk.chinamobile.com/tc/home-family/broadband, 随便查一个地址;
    2. F12 网络面板找到 searchAddress 与 getAddressDetail 两条请求, 分别 Copy as cURL;
    3. 存成一个文件(两条都放进去即可), 例如 config/cmhk.curl;
    4. python3 src/fetch_recog_y.py --mode live --curl-file config/cmhk.curl

    接口带反爬令牌(URL 参数 XGiOG2f705 + Cookie zA7uZWGUB1), 会过期; 脚本遇到
    401/403 会明确提示重新抓取 curl, 不会静默产出空结果.

产出(默认 output/recog_y/):
    recog_y_result.xlsx   四个表: recogY地址 / 楼层房间展开 / 未命中(bad case) / 查询汇总
    recog_y_result.csv    同 recogY地址 表
    recog_y_result.json   完整结构化结果
    recog_y_report.md     人读版报告
    raw/                  每次调用的原始响应(live 模式), 便于核对接口结构
"""

import argparse
import csv
import json
import os
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import cmhk_api
import curl_parser
import hk_address as hk
import screen_bad_cases as sc

Y_FLAG = "Y"


# --------------------------------------------------------------------------

def build_keyword(raw, mode):
    """决定送给接口的查询词."""
    norm, _ = hk.normalize(raw)
    if mode == "raw":
        return raw.strip()
    if mode == "normalized":
        return norm
    if mode == "canonical":
        return hk.canonical(hk.parse(norm)) or norm
    raise ValueError("未知 keyword-mode: %s" % mode)


def offline_verdict(idx, raw):
    """复用离线筛查规则, 给出该地址的静态结论(用于和接口结果交叉验证)."""
    rec = sc.Record(idx, raw)
    sc.check_row_rules(rec)
    return rec


class AddressResult:
    def __init__(self, idx, raw, keyword):
        self.idx = idx
        self.raw = raw
        self.keyword = keyword
        self.candidates = []        # 全部候选
        self.y_items = []           # recogFlag=Y 的候选 + 楼层房间
        self.error = ""
        self.offline = ""           # 离线筛查命中的规则
        self.queried = False        # 是否真的调过接口(plan 模式为 False)

    @property
    def flag_dist(self):
        c = Counter((cand.get("recogFlag") or "-") for cand in self.candidates)
        return ", ".join("%s×%d" % kv for kv in sorted(c.items()))

    @property
    def hit(self):
        return bool(self.y_items)

    @property
    def reason(self):
        if not self.queried:
            return "未调用接口(plan 模式只生成查询计划)"
        if self.error:
            return "接口调用失败: %s" % self.error
        if not self.candidates:
            return "接口无任何候选地址(查无此址)"
        return "有 %d 条候选但无一条 recogFlag=Y(%s)" % (len(self.candidates), self.flag_dist)


def process(records, client, fetch_detail=True):
    """逐条查询并提取; 令牌失效时立即中止, 已完成的结果照常返回."""
    aborted = ""
    for res in records:
        try:
            resp = client.search_address(res.keyword, tag="%02d" % res.idx)
            res.queried = True
            res.candidates = cmhk_api.extract_candidates(resp)
        except cmhk_api.TokenExpired as exc:
            aborted = str(exc)
            break
        except cmhk_api.ApiError as exc:
            res.error = str(exc)
            continue

        for cand in res.candidates:
            if (cand.get("recogFlag") or "").strip().upper() != Y_FLAG:
                continue
            item = {
                "candidate": cand,
                "value": cand.get("value") or "",
                "description": cmhk_api.plain_description(cand),
                "buildingCode": cmhk_api.building_code(cand),
                "ofcaCode": str(cand.get("ofcaCode") or ""),
                "clientType": cand.get("clientType") or "",
                "carriers": ",".join(
                    str(c.get("carrier")) for c in (cand.get("carrierInfo") or [])
                    if isinstance(c, dict) and c.get("carrier")),
                "floors": [], "flats": [], "pairs": [], "detail_status": "未查询",
            }
            if fetch_detail:
                try:
                    detail = client.get_address_detail(cand, tag="%02d_%s" % (res.idx, item["buildingCode"]))
                    floors, flats, pairs = cmhk_api.extract_floors_flats(detail)
                    item.update(floors=floors, flats=flats, pairs=pairs)
                    item["detail_status"] = "成功" if (floors or flats) else "成功但无楼层/房间数据"
                except cmhk_api.TokenExpired as exc:
                    item["detail_status"] = "令牌失效"
                    aborted = str(exc)
                except cmhk_api.ApiError as exc:
                    item["detail_status"] = "失败: %s" % exc
            res.y_items.append(item)
            if aborted:
                break
        if aborted:
            break
    return aborted


def expand_pairs(item):
    """展开成 (楼层, 房间号, 组合来源) 列表."""
    if item["pairs"]:
        return [(f, r, "接口(按层给出房间)") for f, r in item["pairs"]]
    if item["floors"] and item["flats"]:
        return [(f, r, "笛卡尔积(推断)") for f in item["floors"] for r in item["flats"]]
    if item["floors"]:
        return [(f, "", "仅有楼层") for f in item["floors"]]
    if item["flats"]:
        return [("", r, "仅有房间号") for r in item["flats"]]
    return []


# --------------------------------------------------------------------------
# 输出
# --------------------------------------------------------------------------

MAIN_HEADERS = ["序号", "输入地址", "查询关键词", "候选总数", "recogFlag分布",
                "匹配地址", "buildingCode", "ofcaCode", "clientType", "运营商",
                "楼层数", "房间数", "楼层", "房间号", "明细状态", "离线筛查"]

EXPAND_HEADERS = ["序号", "输入地址", "匹配地址", "buildingCode", "楼层", "房间号", "组合来源"]

MISS_HEADERS = ["序号", "输入地址", "查询关键词", "候选总数", "recogFlag分布",
                "未命中原因", "离线筛查"]


def main_rows(records):
    rows = []
    for res in records:
        for item in res.y_items:
            rows.append([
                res.idx, res.raw, res.keyword, len(res.candidates), res.flag_dist,
                item["value"], item["buildingCode"], item["ofcaCode"],
                item["clientType"], item["carriers"],
                len(item["floors"]), len(item["flats"]),
                ",".join(item["floors"]), ",".join(item["flats"]),
                item["detail_status"], res.offline,
            ])
    return rows


def expand_rows(records):
    rows = []
    for res in records:
        for item in res.y_items:
            for floor, flat, src in expand_pairs(item):
                rows.append([res.idx, res.raw, item["value"], item["buildingCode"],
                             floor, flat, src])
    return rows


def miss_rows(records):
    return [[res.idx, res.raw, res.keyword, len(res.candidates), res.flag_dist,
             res.reason, res.offline]
            for res in records if not res.hit]


def write_csv(rows, headers, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)


def write_xlsx(records, summary, path, miss_sheet_title="未命中(bad case)"):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[warn] 未安装 openpyxl, 跳过 xlsx (pip install openpyxl)")
        return False

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        return ws

    wb.remove(wb.active)
    sheet("recogY地址", MAIN_HEADERS, main_rows(records),
          {"输入地址": 30, "查询关键词": 30, "匹配地址": 34, "recogFlag分布": 14,
           "buildingCode": 14, "楼层": 30, "房间号": 24, "明细状态": 18, "离线筛查": 26,
           "运营商": 22, "clientType": 12})
    sheet("楼层房间展开", EXPAND_HEADERS, expand_rows(records),
          {"输入地址": 30, "匹配地址": 34, "buildingCode": 14, "组合来源": 20})
    ws = sheet(miss_sheet_title, MISS_HEADERS, miss_rows(records),
               {"输入地址": 30, "查询关键词": 30, "未命中原因": 46, "离线筛查": 26,
                "recogFlag分布": 14})
    warn = PatternFill("solid", fgColor="FFC7CE")
    for row in ws.iter_rows(min_row=2):
        row[1].fill = warn
    sheet("查询汇总", ["项目", "值"], summary, {"项目": 26, "值": 76})
    wb.save(path)
    return True


def write_report(records, summary, path, mode):
    hit = [r for r in records if r.hit]
    miss = [r for r in records if not r.hit]
    y_total = sum(len(r.y_items) for r in records)
    L = ["# recogFlag=Y 地址 / 楼层 / 房间号 提取报告\n"]
    if mode != "live":
        L.append("> **本次为 `%s` 模式, 不是真实接口数据.** live 模式跑法见 USAGE.md.\n" % mode)
    for k, v in summary:
        L.append("- **%s**: %s" % (k, v))

    L.append("\n## 一、recogFlag=Y 的地址(含楼层与房间号)\n")
    if hit:
        L.append("| 序号 | 输入地址 | 匹配地址 | buildingCode | 楼层数 | 房间数 | 楼层 | 房间号 |")
        L.append("|---|---|---|---|---|---|---|---|")
        for res in hit:
            for item in res.y_items:
                L.append("| %d | %s | %s | %s | %d | %d | %s | %s |" % (
                    res.idx, res.raw, item["value"], item["buildingCode"],
                    len(item["floors"]), len(item["flats"]),
                    ",".join(item["floors"])[:60], ",".join(item["flats"])[:40]))
    else:
        L.append("无\n")

    if mode == "plan":
        L.append("\n## 二、待查询地址清单\n")
        L.append("本次未调用接口, 下表是将要查询的地址与关键词.\n")
    else:
        L.append("\n## 二、未命中的地址(bad case 候选)\n")
        L.append("接口没有返回任何 recogFlag=Y 的候选 —— 这些就是需要人工复核的 bad case.\n")
    if miss:
        L.append("| 序号 | 输入地址 | 候选数 | 未命中原因 | 离线筛查 |")
        L.append("|---|---|---|---|---|")
        for res in miss:
            L.append("| %d | %s | %d | %s | %s |" % (
                res.idx, res.raw, len(res.candidates), res.reason, res.offline or "-"))
    else:
        L.append("无\n")

    L.append("\n## 三、说明\n")
    L.append("- `楼层房间展开` 表把每个 Y 地址拆成 楼层×房间 的组合, 可直接当作后续 "
             "getInstallInfo 的测试用例输入.")
    L.append("- 组合来源标 `笛卡尔积(推断)` 的, 是接口把楼层与房间分成两个平铺列表返回, "
             "脚本无法确定哪层有哪些房间 —— 这类组合需要实际调用 getInstallInfo 验证.")
    L.append("- `离线筛查` 列是不依赖接口的静态规则结论(见 bad_case_report.md), "
             "用来和接口结果交叉印证: 两边都判为问题的最值得优先看.\n")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


# --------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="提取 recogFlag=Y 的地址及楼层/房间号")
    ap.add_argument("--input", default="data/addresses.txt")
    ap.add_argument("--outdir", default="output/recog_y")
    ap.add_argument("--mode", choices=["live", "mock", "plan"], default="mock")
    ap.add_argument("--curl-file", help="live 模式: 含 searchAddress 与 getAddressDetail 的 curl 文件")
    ap.add_argument("--fixture-dir", default="tests/fixtures", help="mock 模式的样例目录")
    ap.add_argument("--keyword-mode", choices=["raw", "normalized", "canonical"],
                    default="raw", help="送给接口的查询词形态, 默认用原始地址")
    ap.add_argument("--delay", type=float, default=1.0, help="两次请求间隔秒数")
    ap.add_argument("--no-detail", action="store_true", help="只查地址, 不查楼层/房间")
    ap.add_argument("--limit", type=int, help="只处理前 N 条, 便于试跑")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8") as f:
        lines = [ln.rstrip("\n") for ln in f if ln.strip()]
    if args.limit:
        lines = lines[: args.limit]

    records = []
    for i, raw in enumerate(lines, 1):
        res = AddressResult(i, raw, build_keyword(raw, args.keyword_mode))
        rec = offline_verdict(i, raw)
        res.offline = ";".join(rec.codes) if rec.codes else "无问题"
        records.append(res)

    os.makedirs(args.outdir, exist_ok=True)
    aborted = ""

    if args.mode == "plan":
        note = "仅生成查询计划, 未调用接口"
    else:
        if args.mode == "live":
            if not args.curl_file:
                ap.error("live 模式必须提供 --curl-file")
            reqs = curl_parser.parse_curl_file(args.curl_file)
            search_tpl = curl_parser.find_by_path(reqs, "searchAddress")
            detail_tpl = curl_parser.find_by_path(reqs, "getAddressDetail")
            if not search_tpl:
                ap.error("%s 里没找到 searchAddress 请求" % args.curl_file)
            if not detail_tpl and not args.no_detail:
                ap.error("%s 里没找到 getAddressDetail 请求(或加 --no-detail)" % args.curl_file)
            client = cmhk_api.CmhkClient(search_tpl, detail_tpl, delay=args.delay,
                                         raw_dir=os.path.join(args.outdir, "raw"))
            note = "真实接口调用"
        else:
            client = cmhk_api.MockClient(args.fixture_dir)
            note = "构造样例(非真实数据), 仅用于验证提取逻辑"
        aborted = process(records, client, fetch_detail=not args.no_detail)

    hit = [r for r in records if r.hit]
    queried = [r for r in records if r.queried]
    y_total = sum(len(r.y_items) for r in records)
    pair_total = len(expand_rows(records))
    summary = [
        ("运行模式", "%s —— %s" % (args.mode, note)),
        ("查询词形态", args.keyword_mode),
        ("输入地址数", str(len(records))),
        ("已查询地址数", "%d / %d" % (len(queried), len(records))),
        ("命中 recogFlag=Y 的地址数", "%d / %d" % (len(hit), len(queried) or len(records))),
        ("recogFlag=Y 候选总条数", str(y_total)),
        ("楼层×房间组合数", str(pair_total)),
        ("未命中(bad case)数",
         str(len(queried) - len(hit)) if queried else "未查询, 不适用"),
        ("调用失败数", str(sum(1 for r in records if r.error))),
    ]
    if aborted:
        summary.append(("中止原因", aborted))

    base = os.path.join(args.outdir, "recog_y_result")
    write_csv(main_rows(records), MAIN_HEADERS, base + ".csv")
    write_csv(expand_rows(records), EXPAND_HEADERS, base + "_expand.csv")
    write_csv(miss_rows(records), MISS_HEADERS, base + "_miss.csv")
    write_xlsx(records, summary, base + ".xlsx",
               "查询计划" if args.mode == "plan" else "未命中(bad case)")
    write_report(records, summary, os.path.join(args.outdir, "recog_y_report.md"), args.mode)
    with open(base + ".json", "w", encoding="utf-8") as f:
        json.dump({
            "summary": dict(summary),
            "records": [{
                "序号": r.idx, "输入地址": r.raw, "查询关键词": r.keyword,
                "候选总数": len(r.candidates), "recogFlag分布": r.flag_dist,
                "离线筛查": r.offline, "错误": r.error,
                "命中": [{k: v for k, v in item.items() if k != "candidate"}
                         for item in r.y_items],
            } for r in records]}, f, ensure_ascii=False, indent=2)

    for k, v in summary:
        print("%s: %s" % (k, v))
    print("已输出到 %s/" % args.outdir)
    if aborted:
        print("\n[中止] %s" % aborted)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
